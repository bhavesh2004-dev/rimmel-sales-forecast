"""
Multi-Brand Client Excel Reporting Module (Phase 5C)
Generates the official simplified 2-sheet client deliverable:
- Sheet 1: Forecast_Inventory (15 exact columns combining sales forecasting & inventory runway)
- Sheet 2: ROP (10 simplified columns with prominent top policy header)
Rule: Exactly 2 sheets in one workbook. AutoFilter and Freeze Panes enabled. Windows file-lock resilient.
"""

import os
import shutil
import logging
from typing import Dict, List, Any, Tuple
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Design System Palette
FONT_TITLE = Font(name='Segoe UI', size=13, bold=True, color='FFFFFF')
FONT_SUBTITLE = Font(name='Segoe UI', size=9, italic=True, color='E0E0E0')
FONT_POLICY = Font(name='Segoe UI', size=10, bold=True, color='1E3A8A')
FONT_HEADER = Font(name='Segoe UI', size=9, bold=True, color='FFFFFF')
FONT_BODY = Font(name='Segoe UI', size=9, color='000000')
FONT_BOLD = Font(name='Segoe UI', size=9, bold=True, color='000000')

FILL_NAVY = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
FILL_STEEL = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid')
FILL_POLICY = PatternFill(start_color='E0E7FF', end_color='E0E7FF', fill_type='solid')
FILL_HEADER = PatternFill(start_color='333F48', end_color='333F48', fill_type='solid')
FILL_ACCENT = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
FILL_ZEBRA = PatternFill(start_color='F9FAFB', end_color='F9FAFB', fill_type='solid')
FILL_WHITE = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
FILL_TOTAL = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')

BORDER_THIN = Border(
    left=Side(style='thin', color='E5E7EB'),
    right=Side(style='thin', color='E5E7EB'),
    top=Side(style='thin', color='E5E7EB'),
    bottom=Side(style='thin', color='E5E7EB')
)

BORDER_POLICY = Border(
    left=Side(style='medium', color='3B82F6'),
    right=Side(style='medium', color='3B82F6'),
    top=Side(style='medium', color='3B82F6'),
    bottom=Side(style='medium', color='3B82F6')
)

BORDER_TOTAL = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='475569'),
    bottom=Side(style='double', color='1E293B')
)

try:
    from .inventory import classify_inventory_risk_and_action, compute_days_of_cover
    from .replenishment import generate_replenishment_layer
except (ImportError, ValueError):
    from src.inventory import classify_inventory_risk_and_action, compute_days_of_cover
    from src.replenishment import generate_replenishment_layer

def generate_forecast_metadata(r: pd.Series, config: Dict[str, Any]) -> Tuple[str, str, str, str]:
    """Generates confidence level, operational risk, recommended action, and human-readable reason."""
    stock = float(r.get('current_stock', 0.0))
    in_stock = int(r.get('in_stock_flag', 1 if stock > 0 else 0))
    cv = float(r.get('cv_30', 0.0))
    fwd = float(r.get('Total Predicted', r.get('10-Day Forecast', 0.0)))
    v30 = float(r.get('v30', 0.0))
    v90 = float(r.get('v90', 0.0))
    
    risk, action = classify_inventory_risk_and_action(stock, in_stock, fwd, v90, cv, config)
    
    if in_stock == 0 or stock <= 0:
        conf = 'LOW'
        reason = 'Inventory availability is limiting observed demand.'
    elif cv > 1.2:
        conf = 'LOW'
        reason = 'Demand has been volatile historically, reducing forecast certainty.'
    elif fwd >= 5.0 and cv <= 0.6:
        conf = 'HIGH'
        reason = 'Stable historical demand with recent momentum remaining consistent.'
    elif fwd < 1.0 and v30 < 0.2:
        conf = 'HIGH'
        reason = 'Consistently low historical sales volume.'
    else:
        conf = 'MEDIUM'
        reason = 'Standard baseline demand run rate.'
        
    return conf, risk, action, reason

def export_client_report(
    df_10d_sku: pd.DataFrame,
    df_rop: pd.DataFrame = None,
    output_path: str = None,
    config: Dict[str, Any] = None
) -> str:
    """
    Exports the simplified 2-sheet client deliverable:
    - Sheet 1: Forecast_Inventory (15 exact columns)
    - Sheet 2: ROP (10 simplified columns with prominent ROP Policy header)
    """
    logger.info(f"Generating simplified 2-sheet client report workbook at: {output_path}...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet
    
    start_date = config['timeline'].get('forward_forecast_start', '2026-09-11')
    end_date = config['timeline'].get('forward_forecast_end', '2026-09-20')
    default_period = f"{start_date} to {end_date}"
    
    # =========================================================================
    # SHEET 1: Forecast_Inventory (Exactly 15 Columns)
    # =========================================================================
    sheet1_name = config.get('reporting', {}).get('forecast_inventory_sheet_name', 'Forecast_Inventory')
    ws1 = wb.create_sheet(sheet1_name)
    ws1.views.sheetView[0].showGridLines = True
    
    headers_s1 = [
        'Brand',
        'SKU',
        'Product',
        'Current Stock',
        'Forecast Period',
        'Amazon Predicted',
        'eBay Predicted',
        'Website Predicted',
        'Other Predicted',
        '10-Day Forecast',
        'Days of Cover',
        'Confidence',
        'Risk',
        'Recommended Action',
        'Reason'
    ]
    
    # Title Banner (Sheet 1)
    last_col_s1 = get_column_letter(len(headers_s1))
    ws1.merge_cells(f'A1:{last_col_s1}1')
    c1_title = ws1['A1']
    c1_title.value = "MULTI-BRAND DEMAND FORECAST & INVENTORY RUNWAY (10-DAY OPERATIONAL PROJECTION)"
    c1_title.font = FONT_TITLE
    c1_title.fill = FILL_NAVY
    c1_title.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws1.row_dimensions[1].height = 28
    
    # Subtitle Banner (Sheet 1)
    ws1.merge_cells(f'A2:{last_col_s1}2')
    c1_sub = ws1['A2']
    c1_sub.value = f"Horizon: {default_period} | Unified Forecast & Dynamic Days of Cover | Additive Channel Law Enforced"
    c1_sub.font = FONT_SUBTITLE
    c1_sub.fill = FILL_STEEL
    c1_sub.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws1.row_dimensions[2].height = 20
    
    # Header Row (Sheet 1, Row 3)
    ws1.row_dimensions[3].height = 24
    for col_idx, h in enumerate(headers_s1, 1):
        cell = ws1.cell(row=3, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_ACCENT if h in ['10-Day Forecast', 'Days of Cover', 'Current Stock'] else FILL_HEADER
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER_THIN
        
    # Populate Sheet 1 Rows
    row_num = 4
    tot_stock = 0
    tot_amz = 0
    tot_ebay = 0
    tot_web = 0
    tot_oth = 0
    tot_fwd = 0
    
    for _, r in df_10d_sku.iterrows():
        brand_raw = r.get('brand_id', r.get('Brand', 'UNKNOWN'))
        brand = 'Rimmel' if 'RIMMEL' in str(brand_raw).upper() else ('Max Factor' if 'MAX' in str(brand_raw).upper() else str(brand_raw))
        sku = r.get('canonical_sku', r.get('SKU', ''))
        prod = r.get('category', '') if pd.notna(r.get('category')) and r.get('category') != '' else r.get('product_title', r.get('Product', sku))
        stock = int(float(r.get('current_stock', r.get('Current Stock', 0.0)))) if pd.notna(r.get('current_stock', r.get('Current Stock', 0.0))) else 0
        period = r.get('Forecast Period', default_period)
        
        amz = int(round(float(r.get('Amazon', r.get('Amazon Predicted', 0.0)))))
        ebay = int(round(float(r.get('eBay', r.get('eBay Predicted', 0.0)))))
        web = int(round(float(r.get('Website', r.get('Website Predicted', 0.0)))))
        oth = int(round(float(r.get('Other', r.get('Other Predicted', 0.0)))))
        
        # Canonical 10-Day Forecast = Sum of individual platform predictions
        tot = amz + ebay + web + oth
        
        doc = compute_days_of_cover(stock, tot)
        conf, risk, action, reason = generate_forecast_metadata(r, config)
        
        tot_stock += stock
        tot_amz += amz
        tot_ebay += ebay
        tot_web += web
        tot_oth += oth
        tot_fwd += tot
        
        row_vals = [
            brand,
            sku,
            prod,
            stock,
            period,
            amz,
            ebay,
            web,
            oth,
            tot,
            doc,
            conf,
            risk,
            action,
            reason
        ]
        
        fill = FILL_ZEBRA if row_num % 2 == 0 else FILL_WHITE
        ws1.row_dimensions[row_num].height = 19
        
        for col_idx, (h, val) in enumerate(zip(headers_s1, row_vals), 1):
            cell = ws1.cell(row=row_num, column=col_idx, value=val)
            cell.font = FONT_BODY
            cell.fill = fill
            cell.border = BORDER_THIN
            
            if h in ['Current Stock', 'Amazon Predicted', 'eBay Predicted', 'Website Predicted', 'Other Predicted', '10-Day Forecast']:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif h == 'Days of Cover':
                if isinstance(val, (int, float)):
                    cell.number_format = '#,##0.0'
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='center', vertical='center')
            elif h in ['Brand', 'Confidence', 'Risk']:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                
        row_num += 1

    # Total Portfolio Summary Row (Sheet 1)
    ws1.row_dimensions[row_num].height = 22
    total_row_s1 = [
        "Total Portfolio", "", "", tot_stock, "",
        tot_amz, tot_ebay, tot_web, tot_oth, tot_fwd,
        round(tot_stock / (tot_fwd / 10.0), 1) if tot_fwd > 0 else "N/A",
        "", "", "", ""
    ]
    for col_idx, (h, val) in enumerate(zip(headers_s1, total_row_s1), 1):
        cell = ws1.cell(row=row_num, column=col_idx, value=val)
        cell.font = FONT_BOLD
        cell.fill = FILL_TOTAL
        cell.border = BORDER_TOTAL
        if isinstance(val, (int, float)):
            cell.number_format = '#,##0.0' if h == 'Days of Cover' else '#,##0'
            cell.alignment = Alignment(horizontal='right', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left' if col_idx == 1 else 'center', vertical='center')

    # Enable AutoFilter and Freeze Panes (Sheet 1)
    ws1.auto_filter.ref = f"A3:{last_col_s1}{row_num - 1}"
    ws1.freeze_panes = 'A4'

    # Auto-fit column widths (Sheet 1)
    for col_idx in range(1, len(headers_s1) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max([len(str(ws1.cell(row=r, column=col_idx).value or '')) for r in range(3, min(row_num + 1, 60))] or [12])
        ws1.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

    # =========================================================================
    # SHEET 2: ROP (Exactly 10 Simplified Columns with ROP Policy Header)
    # =========================================================================
    sheet2_name = config.get('reporting', {}).get('rop_sheet_name', 'ROP')
    ws2 = wb.create_sheet(sheet2_name)
    ws2.views.sheetView[0].showGridLines = True
    
    headers_s2 = [
        'Brand',
        'SKU',
        'Product',
        'Current Stock',
        '10-Day Forecast',
        'Avg Daily Usage',
        'Lead-Time Demand',
        'Target Stock',
        'Replenishment Qty',
        'ROP Status'
    ]
    last_col_s2 = get_column_letter(len(headers_s2))
    
    # 1. Title Banner (Row 1)
    ws2.merge_cells(f'A1:{last_col_s2}1')
    c2_title = ws2['A1']
    c2_title.value = "MULTI-BRAND REORDER POINT (ROP) & REPLENISHMENT PLANNING"
    c2_title.font = FONT_TITLE
    c2_title.fill = FILL_NAVY
    c2_title.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws2.row_dimensions[1].height = 28
    
    # 2. Compact ROP Policy Header Section (Row 2)
    rop_cfg = config.get('replenishment', {})
    lt = rop_cfg.get('lead_time_days', 10)
    min_stk = rop_cfg.get('minimum_stock_level', 6)
    
    ws2.merge_cells(f'A2:{last_col_s2}2')
    c2_policy = ws2['A2']
    c2_policy.value = f"ROP POLICY: Lead Time = {lt} Days | Minimum Stock Level = {min_stk} Units | Target Stock = (Avg Daily Usage × 10) + 6"
    c2_policy.font = FONT_POLICY
    c2_policy.fill = FILL_POLICY
    c2_policy.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    c2_policy.border = BORDER_POLICY
    ws2.row_dimensions[2].height = 22
    
    # 3. Blank Row (Row 3)
    ws2.row_dimensions[3].height = 10
    
    # 4. Table Header Row (Row 4)
    ws2.row_dimensions[4].height = 24
    for col_idx, h in enumerate(headers_s2, 1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.font = FONT_HEADER
        cell.fill = FILL_ACCENT if h in ['Target Stock', 'Replenishment Qty', 'ROP Status'] else FILL_HEADER
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = BORDER_THIN
        
    # Generate ROP dataframe if not supplied
    if df_rop is None:
        df_rop = generate_replenishment_layer(df_10d_sku, config)
        
    # 5. Populate Data Rows (Row 5 onwards)
    row_num2 = 5
    tot_rop_stock = 0
    tot_rop_fwd = 0
    tot_rop_ltd = 0.0
    tot_rop_target = 0.0
    tot_rop_qty = 0
    
    for _, r in df_rop.iterrows():
        b = r['Brand']
        sku = r['SKU']
        prod = r['Product']
        stk = int(r['Current Stock'])
        fwd = int(r['10-Day Forecast'])
        adu = float(r['Avg Daily Usage'])
        ltd = float(r['Lead-Time Demand'])
        tgt = float(r['Target Stock'])
        rep_qty = int(r['Replenishment Qty'])
        status = r['ROP Status']
        
        tot_rop_stock += stk
        tot_rop_fwd += fwd
        tot_rop_ltd += ltd
        tot_rop_target += tgt
        tot_rop_qty += rep_qty
        
        row_vals2 = [
            b, sku, prod, stk, fwd, adu, ltd, tgt, rep_qty, status
        ]
        
        fill2 = FILL_ZEBRA if row_num2 % 2 == 0 else FILL_WHITE
        ws2.row_dimensions[row_num2].height = 19
        
        for col_idx, (h, val) in enumerate(zip(headers_s2, row_vals2), 1):
            cell = ws2.cell(row=row_num2, column=col_idx, value=val)
            cell.font = FONT_BODY
            cell.fill = fill2
            cell.border = BORDER_THIN
            
            if h in ['Current Stock', '10-Day Forecast', 'Lead-Time Demand', 'Target Stock', 'Replenishment Qty']:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif h == 'Avg Daily Usage':
                cell.number_format = '#,##0.0'
                cell.alignment = Alignment(horizontal='right', vertical='center')
            elif h in ['Brand', 'ROP Status']:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                
        row_num2 += 1

    # 6. Total Portfolio Summary Row (Sheet 2)
    ws2.row_dimensions[row_num2].height = 22
    total_row_s2 = [
        "Total Portfolio", "", "", tot_rop_stock, tot_rop_fwd,
        round(tot_rop_fwd / 10.0, 1), round(tot_rop_ltd), round(tot_rop_target), tot_rop_qty, ""
    ]
    for col_idx, (h, val) in enumerate(zip(headers_s2, total_row_s2), 1):
        cell = ws2.cell(row=row_num2, column=col_idx, value=val)
        cell.font = FONT_BOLD
        cell.fill = FILL_TOTAL
        cell.border = BORDER_TOTAL
        if isinstance(val, (int, float)):
            cell.number_format = '#,##0.0' if h == 'Avg Daily Usage' else '#,##0'
            cell.alignment = Alignment(horizontal='right', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left' if col_idx == 1 else 'center', vertical='center')

    # Enable AutoFilter and Freeze Panes (Sheet 2)
    ws2.auto_filter.ref = f"A4:{last_col_s2}{row_num2 - 1}"
    ws2.freeze_panes = 'A5'

    # Auto-fit column widths (Sheet 2)
    for col_idx in range(1, len(headers_s2) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max([len(str(ws2.cell(row=r, column=col_idx).value or '')) for r in range(4, min(row_num2 + 1, 60))] or [12])
        ws2.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

    # =========================================================================
    # Save Workbook with Windows File Lock Resilience
    # =========================================================================
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    candidate_paths = [
        output_path,
        output_path.replace('.xlsx', '_UPDATED.xlsx'),
        output_path.replace('.xlsx', '_LATEST.xlsx'),
        output_path.replace('.xlsx', '_FINAL.xlsx')
    ]
    for p in candidate_paths:
        try:
            wb.save(p)
            logger.info(f"Official client report successfully saved to: {p}")
            return p
        except PermissionError:
            logger.warning(f"File lock detected on {p}. Trying next candidate...")
            
    import datetime
    ts = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    fallback_path = output_path.replace('.xlsx', f'_{ts}.xlsx')
    wb.save(fallback_path)
    logger.info(f"Saved to timestamped fallback: {fallback_path}")
    return fallback_path
