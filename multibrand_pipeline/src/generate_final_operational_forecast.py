"""
Final Operational Model Refit & 10-Day Operational Forecast Generator (Phase 5C)
Generates the official simplified client workbook:
RIMMEL_MAX_FACTOR_10DAY_OPERATIONAL_FORECAST.xlsx

Workbook Architecture (Exactly 2 Sheets):
├── Sheet 1: Forecast_Inventory (Exactly 15 Columns)
│   ├── Brand
│   ├── SKU
│   ├── Product
│   ├── Current Stock
│   ├── Forecast Period
│   ├── Amazon Predicted
│   ├── eBay Predicted
│   ├── Website Predicted
│   ├── Other Predicted
│   ├── 10-Day Forecast
│   ├── Days of Cover
│   ├── Confidence
│   ├── Risk
│   ├── Recommended Action
│   └── Reason
│
└── Sheet 2: ROP (Exactly 10 Simplified Columns with ROP Policy Header)
    ├── Brand
    ├── SKU
    ├── Product
    ├── Current Stock
    ├── 10-Day Forecast
    ├── Avg Daily Usage
    ├── Lead-Time Demand
    ├── Target Stock
    ├── Replenishment Qty
    └── ROP Status
"""

import os
import sys
import json
import yaml
import shutil
import logging
import pandas as pd
import numpy as np

# Ensure multibrand_pipeline root and src are on path
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PIPELINE_ROOT = os.path.dirname(CURRENT_DIR)
WORKSPACE_ROOT = os.path.dirname(PIPELINE_ROOT)

if PIPELINE_ROOT not in sys.path:
    sys.path.insert(0, PIPELINE_ROOT)
if CURRENT_DIR not in sys.path:
    sys.path.insert(0, CURRENT_DIR)

from src.train import train_and_evaluate
from src.forecast import generate_forward_simulation
from src.replenishment import generate_replenishment_layer, validate_replenishment_invariants
from src.reporting import export_client_report

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
logger = logging.getLogger("final_operational_forecast")

def main():
    logger.info("==================================================================")
    logger.info("STARTING FINAL OPERATIONAL REFIT & 10-DAY OPERATIONAL FORECAST")
    logger.info("==================================================================")

    config_path = os.path.join(PIPELINE_ROOT, "config", "pipeline_config.yaml")
    schema_path = os.path.join(PIPELINE_ROOT, "config", "feature_schema.json")
    parquet_path = os.path.join(WORKSPACE_ROOT, "experiments", "global_multibrand_lgbm", "data", "global_training_dataset.parquet")

    with open(config_path, "r", encoding="utf-8") as f:
        config = yaml.safe_load(f)
    with open(schema_path, "r", encoding="utf-8") as f:
        schema = json.load(f)

    logger.info(f"Loaded config from: {config_path}")
    logger.info(f"Loaded feature schema: {schema['schema_version']} ({len(schema['feature_list'])} features)")

    logger.info(f"Loading multi-brand dataset from: {parquet_path}...")
    df = pd.read_parquet(parquet_path)
    logger.info(f"Dataset loaded: {len(df):,} rows spanning {df['date'].min()} to {df['date'].max()}.")

    # 1. Operational Model Refit (Train 2025-08-01 -> 2026-09-10)
    logger.info("\n[1/4] Refitting Global LightGBM on complete validated timeline (2025-08-01 to 2026-09-10)...")
    model, model_meta = train_and_evaluate(
        df_features=df,
        config=config,
        feature_schema=schema,
        refit_mode=True
    )
    logger.info("Operational model successfully refit and serialized.")

    # 2. Forward Day-by-Day 10-Day Simulation (2026-09-11 -> 2026-09-20)
    logger.info(f"\n[2/4] Generating day-by-day forward simulation ({config['timeline']['forward_forecast_start']} to {config['timeline']['forward_forecast_end']})...")
    df_daily_fwd, df_10d_sku = generate_forward_simulation(
        model=model,
        df_features=df,
        config=config,
        feature_schema=schema
    )
    logger.info(f"Forward simulation generated {len(df_10d_sku):,} aggregated SKU records.")

    # 3. Downstream ROP Replenishment Layer Generation & Mathematical Assertions
    logger.info("\n[3/4] Generating downstream ROP replenishment layer & running mathematical validations...")
    df_rop = generate_replenishment_layer(df_10d_sku, config)
    validate_replenishment_invariants(df_rop, config)

    # 4. Export Official Workbook: Exactly 2 sheets (Forecast_Inventory & ROP)
    output_filename = "RIMMEL_MAX_FACTOR_10DAY_OPERATIONAL_FORECAST.xlsx"
    target_path_root = os.path.join(WORKSPACE_ROOT, output_filename)
    target_path_reports = os.path.join(PIPELINE_ROOT, "reports", output_filename)

    logger.info(f"\n[4/4] Exporting official 2-sheet Excel workbook to: {target_path_root}...")
    saved_path = export_client_report(
        df_10d_sku=df_10d_sku,
        df_rop=df_rop,
        output_path=target_path_root,
        config=config
    )

    # Also save an archived copy in multibrand_pipeline/reports/
    try:
        shutil.copyfile(saved_path, target_path_reports)
        logger.info(f"Archived copy saved to: {target_path_reports}")
    except Exception as e:
        logger.warning(f"Could not copy to {target_path_reports}: {e}")

    # 5. Comprehensive Quality & Integrity Audit
    logger.info("\n==================================================================")
    logger.info("AUDITING GENERATED WORKBOOK INTEGRITY (PHASE 5C)")
    logger.info("==================================================================")
    import openpyxl
    wb = openpyxl.load_workbook(saved_path, data_only=True)
    sheet_names = wb.sheetnames
    logger.info(f"Sheet names present: {sheet_names}")
    assert len(sheet_names) == 2, f"Expected exactly 2 sheets, got {len(sheet_names)}: {sheet_names}"
    assert sheet_names == ["Forecast_Inventory", "ROP"], f"Expected ['Forecast_Inventory', 'ROP'], got: {sheet_names}"

    # Audit Sheet 1: Forecast_Inventory
    ws1 = wb["Forecast_Inventory"]
    headers1 = [ws1.cell(3, c).value for c in range(1, 16)]
    logger.info(f"Sheet 1 Headers (15 columns): {headers1}")
    expected_headers1 = [
        'Brand', 'SKU', 'Product', 'Current Stock', 'Forecast Period',
        'Amazon Predicted', 'eBay Predicted', 'Website Predicted', 'Other Predicted',
        '10-Day Forecast', 'Days of Cover', 'Confidence', 'Risk', 'Recommended Action', 'Reason'
    ]
    assert headers1 == expected_headers1, f"Sheet 1 Headers mismatch! Got: {headers1}"
    data_rows1 = ws1.max_row - 4
    assert data_rows1 == 1108, f"Expected 1,108 SKUs in Sheet 1, got {data_rows1}"
    assert ws1.auto_filter.ref is not None, "AutoFilter missing on Sheet 1!"

    # Audit Sheet 2: ROP
    ws2 = wb["ROP"]
    policy_banner = ws2.cell(2, 1).value
    logger.info(f"Sheet 2 Policy Banner: {policy_banner}")
    assert "Lead Time = 10 Days" in str(policy_banner), "Lead Time missing from ROP Policy banner!"
    assert "Minimum Stock Level = 6 Units" in str(policy_banner), "Minimum Stock Level missing from ROP Policy banner!"

    headers2 = [ws2.cell(4, c).value for c in range(1, 11)]
    logger.info(f"Sheet 2 Headers (10 columns): {headers2}")
    expected_headers2 = [
        'Brand', 'SKU', 'Product', 'Current Stock', '10-Day Forecast',
        'Avg Daily Usage', 'Lead-Time Demand', 'Target Stock',
        'Replenishment Qty', 'ROP Status'
    ]
    assert headers2 == expected_headers2, f"Sheet 2 Headers mismatch! Got: {headers2}\nExpected: {expected_headers2}"
    data_rows2 = ws2.max_row - 5
    assert data_rows2 == 1108, f"Expected 1,108 SKUs in Sheet 2, got {data_rows2}"
    assert ws2.auto_filter.ref is not None, "AutoFilter missing on Sheet 2!"

    # Extract Data Rows for forensic checks
    df_s1 = pd.DataFrame(ws1.iter_rows(min_row=4, max_row=1111, values_only=True), columns=headers1)
    df_s2 = pd.DataFrame(ws2.iter_rows(min_row=5, max_row=1112, values_only=True), columns=headers2)

    # Audit Additive Channel Law
    channel_sum = df_s1['Amazon Predicted'] + df_s1['eBay Predicted'] + df_s1['Website Predicted'] + df_s1['Other Predicted']
    additive_violations = (channel_sum != df_s1['10-Day Forecast']).sum()
    logger.info(f"Additive Channel Law Violations: {additive_violations}")
    assert additive_violations == 0, f"Found {additive_violations} additive channel violations!"

    # Check for 999 sentinels in both sheets
    sentinel_s1 = (df_s1.astype(str) == '999').sum().sum() + (df_s1.astype(str) == '999.0').sum().sum()
    sentinel_s2 = (df_s2.astype(str) == '999').sum().sum() + (df_s2.astype(str) == '999.0').sum().sum()
    logger.info(f"Sentinel '999' counts: Sheet 1 = {sentinel_s1}, Sheet 2 = {sentinel_s2}")
    assert sentinel_s1 == 0 and sentinel_s2 == 0, "Found '999' sentinels in workbook!"

    # Section 15 Invariant: Projected Stock after LTD and replenishment >= 6
    proj_stock = df_s2['Current Stock'] + df_s2['Replenishment Qty'] - df_s2['Lead-Time Demand']
    min_proj = proj_stock.min()
    logger.info(f"Minimum Projected Stock after LTD and replenishment: {min_proj:.2f} units (Expected >= 6.00)")
    assert min_proj >= 6.0 - 1e-4, f"Projected stock dropped below 6 units! Min was: {min_proj}"

    # Summary Stats
    print("\n==================================================================")
    print("OPERATIONAL 10-DAY FORECAST & ROP SUMMARY (2026-09-11 to 2026-09-20)")
    print("==================================================================")
    print(f"Total Canonical SKUs: {len(df_s1):,}")
    print(f"Total Current Warehouse Stock: {df_s1['Current Stock'].sum():,} units")
    print(f"Total 10-Day Forecast Across All Channels: {df_s1['10-Day Forecast'].sum():,} units")
    print(f"  Amazon:  {df_s1['Amazon Predicted'].sum():,} units ({df_s1['Amazon Predicted'].sum() / df_s1['10-Day Forecast'].sum() * 100:.1f}%)")
    print(f"  eBay:    {df_s1['eBay Predicted'].sum():,} units ({df_s1['eBay Predicted'].sum() / df_s1['10-Day Forecast'].sum() * 100:.1f}%)")
    print(f"  Website: {df_s1['Website Predicted'].sum():,} units ({df_s1['Website Predicted'].sum() / df_s1['10-Day Forecast'].sum() * 100:.1f}%)")
    print(f"  Other:   {df_s1['Other Predicted'].sum():,} units ({df_s1['Other Predicted'].sum() / df_s1['10-Day Forecast'].sum() * 100:.1f}%)")
    
    print("\nBrand Breakdown:")
    for b in df_s1['Brand'].unique():
        b_sub1 = df_s1[df_s1['Brand'] == b]
        b_sub2 = df_s2[df_s2['Brand'] == b]
        print(f"  {b}: {b_sub1['10-Day Forecast'].sum():,} units projected | {b_sub2['Replenishment Qty'].sum():,} units recommended order across {len(b_sub1):,} SKUs")

    print("\nROP Status Distribution:")
    for stat, cnt in df_s2['ROP Status'].value_counts().items():
        print(f"  {stat}: {cnt:,} SKUs ({cnt / len(df_s2) * 100:.1f}%)")

    print("\nTotal Recommended Replenishment Quantity Across Portfolio:")
    print(f"  Total Units to Order: {df_s2['Replenishment Qty'].sum():,} units")
    skus_ordering = (df_s2['Replenishment Qty'] > 0).sum()
    print(f"  SKUs Requiring Purchase Order: {skus_ordering:,} / {len(df_s2):,} ({skus_ordering / len(df_s2) * 100:.1f}%)")

    print("\nWorkbook successfully generated and certified at:")
    print(f"-> {saved_path}")
    print(f"-> {target_path_reports}")

if __name__ == "__main__":
    main()
