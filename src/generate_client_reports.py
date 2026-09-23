"""
Client Report Generation Pipeline: SKU-Level 10-Day Summary
===========================================================
Rimmel Multi-Platform Demand Forecasting System

Produces:
1. reports/Rimmel_Validation_Sep01_Sep10_2026.xlsx
   - Sheet 1: 'SKU Validation Summary' (ONE ROW PER CANONICAL SKU, 674 rows)
   - Sheet 2: 'Platform Summary'
2. reports/Rimmel_Forward_Forecast_Sep11_Sep20_2026.xlsx
   - Sheet 1: 'SKU Forecast Summary' (ONE ROW PER CANONICAL SKU, 674 rows)
   - Sheet 2: 'Platform Summary'
   - Sheet 3: 'Inventory Actions' (Shared warehouse stock strictly unsummed)
3. Backward-compatible copies:
   - reports/validation_report_sep_01_to_10_2026.xlsx
   - reports/production_forecast_sep_11_to_20_2026.xlsx
4. Cache deliverables in data/processed/ and reports/
"""

import os
import sys
import time
import json
import shutil
import pickle
import sqlite3
import numpy as np
import pandas as pd
import lightgbm as lgb
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
if BASE_DIR not in sys.path:
    sys.path.append(BASE_DIR)

# Formatting Constants
FONT_NAME = 'Calibri'
NAVY_FILL = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
AMZ_FILL = PatternFill(start_color='C0392B', end_color='C0392B', fill_type='solid')
EBAY_FILL = PatternFill(start_color='2980B9', end_color='2980B9', fill_type='solid')
WEB_FILL = PatternFill(start_color='27AE60', end_color='27AE60', fill_type='solid')
OTH_FILL = PatternFill(start_color='8E44AD', end_color='8E44AD', fill_type='solid')
TOT_FILL = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
CHARCOAL_FILL = PatternFill(start_color='34495E', end_color='34495E', fill_type='solid')
LIGHT_BLUE_FILL = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color='FFFFFF')
BOLD_FONT = Font(name=FONT_NAME, size=11, bold=True)
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color='1F4E78')
SUBTITLE_FONT = Font(name=FONT_NAME, size=11, italic=True, color='595959')
REGULAR_FONT = Font(name=FONT_NAME, size=11)

THIN_SIDE = Side(border_style='thin', color='D9D9D9')
BORDER_ALL = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
BORDER_TOP_THIN_BOTTOM_DOUBLE = Border(
    top=Side(border_style='thin', color='000000'),
    bottom=Side(border_style='double', color='000000')
)

def style_header_row_custom(ws, headers, fills):
    ws.row_dimensions[1].height = 28
    for col_idx, (header_name, fill_style) in enumerate(zip(headers, fills), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_name)
        cell.font = HEADER_FONT
        cell.fill = fill_style
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER_ALL

def auto_fit_columns_custom(ws, max_cols, reason_col_idx=None, min_width=12):
    for col in range(1, max_cols + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(1, min(ws.max_row + 1, 700)):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        if reason_col_idx and col == reason_col_idx:
            ws.column_dimensions[col_letter].width = max(max_len + 4, 48)
        else:
            ws.column_dimensions[col_letter].width = max(max_len + 4, min_width)

def generate_sku_reason(stock, in_stock, v14, v30, v90, cv_30, total_act, total_pred):
    """
    Produces concise, client-friendly business explanations.
    Strictly grounded in verified historical data and model signals.
    """
    if in_stock == 0 or (stock is not None and not pd.isna(stock) and stock <= 0):
        return "Inventory availability is limiting observed demand."
    if cv_30 is not None and not pd.isna(cv_30) and cv_30 > 1.2:
        return "Demand has been volatile historically, reducing forecast certainty."
    if (v30 is None or pd.isna(v30) or v30 <= 0.05) and total_pred <= 0:
        return "Demand is intermittent with limited recent sales evidence."
    if v30 is not None and v14 is not None and not pd.isna(v30) and not pd.isna(v14) and v30 > 0.1:
        if v14 > 1.25 * v30 and total_pred > 0:
            return "Recent demand is increasing relative to the longer-term baseline."
        if v14 < 0.75 * v30:
            return "Demand is declining compared with the recent historical period."
    if v30 is not None and not pd.isna(v30) and v30 < 0.2 and total_pred > 0:
        return "Low historical demand makes the forecast less certain."
    return "Stable historical demand with recent demand remaining consistent."

def classify_sku_confidence_risk(stock, in_stock, v14, v30, cv_30, total_pred):
    """
    Assigns Confidence (HIGH/MEDIUM/LOW) and Risk (NORMAL/STOCKOUT RISK/HIGH VOLATILITY/LOW DEMAND).
    """
    if in_stock == 0 or (stock is not None and not pd.isna(stock) and stock <= 0):
        conf = "LOW"
        risk = "STOCKOUT RISK"
    elif cv_30 is not None and not pd.isna(cv_30) and cv_30 > 1.2:
        conf = "LOW"
        risk = "HIGH VOLATILITY"
    elif (v30 is None or pd.isna(v30) or v30 <= 0.05) and total_pred <= 0:
        conf = "HIGH"
        risk = "LOW DEMAND"
    elif v30 is not None and v14 is not None and not pd.isna(v30) and not pd.isna(v14) and v30 > 0.1 and (v14 > 1.25 * v30 or v14 < 0.75 * v30):
        conf = "MEDIUM"
        risk = "NORMAL"
    else:
        conf = "HIGH"
        risk = "NORMAL"
    return conf, risk

def get_recommended_action(stock, in_stock, cv_30, total_pred):
    """
    Assigns Operational Recommended Action for Inventory Planning.
    """
    if in_stock == 0 or (stock is not None and not pd.isna(stock) and stock <= 0):
        return "Urgent Restock"
    if stock is not None and not pd.isna(stock) and stock < total_pred:
        return "Reorder Required"
    if cv_30 is not None and not pd.isna(cv_30) and cv_30 > 1.2:
        return "Monitor Closely"
    if total_pred <= 0:
        return "No Action Needed"
    return "Maintain Stock"

def main():
    db_path = os.path.join(BASE_DIR, 'data', 'rimmel_clean.db')
    reports_dir = os.path.join(BASE_DIR, 'reports')
    processed_dir = os.path.join(BASE_DIR, 'data', 'processed')
    models_dir = os.path.join(BASE_DIR, 'models')
    os.makedirs(reports_dir, exist_ok=True)
    os.makedirs(processed_dir, exist_ok=True)
    os.makedirs(models_dir, exist_ok=True)

    print("=" * 80)
    print("RIMMEL DEMAND FORECASTING: REBUILD CLIENT REPORTS (SKU-LEVEL 10-DAY SUMMARY)")
    print(f"Timestamp: {time.strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"Database: {db_path}")
    print("=" * 80)

    # 1. LOAD DATA & FEATURES
    print("\n[STEP 1] Loading data from SQLite database...")
    conn = sqlite3.connect(db_path)
    df = pd.read_sql_query("SELECT * FROM ml_features_zero", conn)
    conn.close()
    print(f"  Loaded {len(df):,} records.")

    features_path = os.path.join(models_dir, 'production_features.json')
    with open(features_path, 'r') as f:
        meta_features = json.load(f)
    feature_cols = meta_features['feature_list']
    cat_cols = meta_features['categorical_features']
    exact_exp6_params = meta_features['exact_hyperparameters']

    for c in cat_cols:
        if c in feature_cols:
            df[c] = df[c].astype('category')

    all_skus = sorted(df['canonical_sku'].unique().tolist())
    num_unique_skus = len(all_skus)
    print(f"  Unique catalog SKUs: {num_unique_skus}")

    # Primary Product/Category lookup per SKU
    sku_meta_df = df.groupby('canonical_sku', observed=True).agg({
        'category': 'first',
        'current_stock': 'last',
        'in_stock_flag': 'last',
        'v7': 'last',
        'v14': 'last',
        'v30': 'last',
        'v90': 'last',
        'cv_30': 'last'
    }).reindex(all_skus).reset_index()

    product_lookup = {}
    for _, r in sku_meta_df.iterrows():
        cat = str(r['category']) if pd.notna(r['category']) else 'Cosmetics General'
        primary_prod = cat.split(',')[0].strip()
        product_lookup[r['canonical_sku']] = primary_prod

    # 2. RETROSPECTIVE VALIDATION PREDICTIONS (SEP 01 - SEP 10, 2026)
    print("\n[STEP 2] Running Retrospective Holdout Validation on Sep 01-10, 2026...")
    train_mask_val = df['date'] < '2026-09-01'
    holdout_mask = (df['date'] >= '2026-09-01') & (df['date'] <= '2026-09-10')

    X_tr_val = df.loc[train_mask_val, feature_cols]
    y_tr_val = df.loc[train_mask_val, 'model_units_sold']

    X_holdout = df.loc[holdout_mask, feature_cols]
    holdout_df = df.loc[holdout_mask].copy()

    val_model = lgb.LGBMRegressor(**exact_exp6_params)
    val_model.fit(X_tr_val, y_tr_val)
    pred_raw_val = np.clip(val_model.predict(X_holdout), 0, None)

    # Combined Calibration (Exp6)
    z_mask_val = (holdout_df['v7'] == 0) & (holdout_df['v14'] == 0) & (holdout_df['v30'] == 0)
    ev_mask_val = (holdout_df['promo_days_30'] > 0) | (holdout_df['amazon_sessions_momentum'] > 1.25)
    calib_z_val = (z_mask_val & (~ev_mask_val)).values
    calib_stk_val = ((holdout_df['in_stock_flag'] == 0) & (holdout_df['has_inventory_signal'] == 1)).values

    exp6_val_preds = pred_raw_val.copy()
    exp6_val_preds[calib_z_val] *= 0.10
    exp6_val_preds[calib_stk_val] *= 0.10

    holdout_df['pred_decimal'] = exp6_val_preds
    holdout_df['act_units'] = holdout_df['observed_units_sold']

    val_tot_act = float(holdout_df['act_units'].sum())
    val_tot_pred_cont = float(holdout_df['pred_decimal'].sum())
    print(f"  Holdout Benchmark: Actual = {val_tot_act:,.1f}, Continuous Pred = {val_tot_pred_cont:,.1f}")

    # 3. FORWARD PRODUCTION FORECAST (SEP 11 - SEP 20, 2026)
    print("\n[STEP 3] Generating Forward Production Forecast for Sep 11-20, 2026...")
    prod_model_path = os.path.join(models_dir, 'production_lgbm_model.pkl')
    if os.path.exists(prod_model_path):
        with open(prod_model_path, 'rb') as f:
            prod_model = pickle.load(f)
    else:
        prod_train_mask = df['date'] <= '2026-09-10'
        prod_model = lgb.LGBMRegressor(**exact_exp6_params)
        prod_model.fit(df.loc[prod_train_mask, feature_cols], df.loc[prod_train_mask, 'model_units_sold'])
        with open(prod_model_path, 'wb') as f:
            pickle.dump(prod_model, f)

    sep10_df = df[df['date'] == '2026-09-10'].copy()
    forward_dates = [
        ('2026-09-11', 4, 0), ('2026-09-12', 5, 1), ('2026-09-13', 6, 1),
        ('2026-09-14', 0, 0), ('2026-09-15', 1, 0), ('2026-09-16', 2, 0),
        ('2026-09-17', 3, 0), ('2026-09-18', 4, 0), ('2026-09-19', 5, 1),
        ('2026-09-20', 6, 1)
    ]

    forward_records = []
    for f_date, dow, is_wknd in forward_dates:
        day_feat = sep10_df[feature_cols].copy()
        day_feat['day_of_week'] = dow
        day_feat['is_weekend'] = is_wknd
        raw_p = np.clip(prod_model.predict(day_feat), 0, None)

        z_mask_f = (sep10_df['v7'] == 0) & (sep10_df['v14'] == 0) & (sep10_df['v30'] == 0)
        ev_mask_f = (sep10_df['promo_days_30'] > 0) | (sep10_df['amazon_sessions_momentum'] > 1.25)
        calib_z_f = (z_mask_f & (~ev_mask_f)).values
        calib_stk_f = ((sep10_df['in_stock_flag'] == 0) & (sep10_df['has_inventory_signal'] == 1)).values

        p = raw_p.copy()
        p[calib_z_f] *= 0.10
        p[calib_stk_f] *= 0.10

        for idx, (_, r) in enumerate(sep10_df.iterrows()):
            forward_records.append({
                'canonical_sku': r['canonical_sku'],
                'platform_group': r['platform_group'],
                'pred': p[idx]
            })

    forward_df = pd.DataFrame(forward_records)
    fwd_tot_pred_cont = float(forward_df['pred'].sum())
    print(f"  Forward Production: Continuous Pred = {fwd_tot_pred_cont:,.1f}")

    # =========================================================================
    # 4. AGGREGATE TO SKU-LEVEL 10-DAY SUMMARY
    # =========================================================================
    print("\n[STEP 4] Aggregating continuous predictions across 10 days for each SKU...")

    # A. VALIDATION 10-DAY AGGREGATION
    val_sku_plat = holdout_df.groupby(['canonical_sku', 'platform_group'], observed=False).agg({
        'act_units': 'sum',
        'pred_decimal': 'sum'
    }).reset_index()

    val_piv_act = val_sku_plat.pivot(index='canonical_sku', columns='platform_group', values='act_units').fillna(0).reindex(all_skus, fill_value=0)
    val_piv_pred_cont = val_sku_plat.pivot(index='canonical_sku', columns='platform_group', values='pred_decimal').fillna(0.0).reindex(all_skus, fill_value=0.0)

    for p in ['Amazon', 'eBay', 'Website', 'Other']:
        if p not in val_piv_act.columns: val_piv_act[p] = 0
        if p not in val_piv_pred_cont.columns: val_piv_pred_cont[p] = 0.0

    # Non-lossy whole unit rounding: round after summing all 10 days
    val_piv_pred_int = np.round(val_piv_pred_cont).astype(int)

    val_amz_act = val_piv_act['Amazon'].astype(int)
    val_amz_pred = val_piv_pred_int['Amazon']
    val_ebay_act = val_piv_act['eBay'].astype(int)
    val_ebay_pred = val_piv_pred_int['eBay']
    val_web_act = val_piv_act['Website'].astype(int)
    val_web_pred = val_piv_pred_int['Website']
    val_oth_act = val_piv_act['Other'].astype(int)
    val_oth_pred = val_piv_pred_int['Other']

    val_tot_act_col = val_amz_act + val_ebay_act + val_web_act + val_oth_act
    val_tot_pred_col = val_amz_pred + val_ebay_pred + val_web_pred + val_oth_pred
    val_variance_col = val_tot_pred_col - val_tot_act_col

    val_conf_list = []
    val_risk_list = []
    val_reason_list = []

    for sku in all_skus:
        r = sku_meta_df[sku_meta_df['canonical_sku'] == sku].iloc[0]
        c_act = val_tot_act_col[sku]
        c_pred = val_tot_pred_col[sku]
        conf, risk = classify_sku_confidence_risk(r['current_stock'], r['in_stock_flag'], r['v14'], r['v30'], r['cv_30'], c_pred)
        reason = generate_sku_reason(r['current_stock'], r['in_stock_flag'], r['v14'], r['v30'], r['v90'], r['cv_30'], c_act, c_pred)
        val_conf_list.append(conf)
        val_risk_list.append(risk)
        val_reason_list.append(reason)

    val_summary_df = pd.DataFrame({
        'SKU': all_skus,
        'Product': [product_lookup[s] for s in all_skus],
        'Validation Period': 'Sep 1–10, 2026',
        'Amazon Actual': val_amz_act.values,
        'Amazon Predicted': val_amz_pred.values,
        'eBay Actual': val_ebay_act.values,
        'eBay Predicted': val_ebay_pred.values,
        'Website Actual': val_web_act.values,
        'Website Predicted': val_web_pred.values,
        'Other Actual': val_oth_act.values,
        'Other Predicted': val_oth_pred.values,
        'Total Actual': val_tot_act_col.values,
        'Total Predicted': val_tot_pred_col.values,
        'Variance': val_variance_col.values,
        'Confidence': val_conf_list,
        'Risk': val_risk_list,
        'Reason': val_reason_list
    })

    print(f"  Validation SKU Summary rows: {len(val_summary_df):,}")
    print(f"    Total Actual Units: {val_summary_df['Total Actual'].sum():,}")
    print(f"    Total Predicted Units (10d rounded sum): {val_summary_df['Total Predicted'].sum():,}")

    # B. FORWARD FORECAST 10-DAY AGGREGATION
    fwd_sku_plat = forward_df.groupby(['canonical_sku', 'platform_group'], observed=False)['pred'].sum().reset_index()
    fwd_piv_pred_cont = fwd_sku_plat.pivot(index='canonical_sku', columns='platform_group', values='pred').fillna(0.0).reindex(all_skus, fill_value=0.0)

    for p in ['Amazon', 'eBay', 'Website', 'Other']:
        if p not in fwd_piv_pred_cont.columns: fwd_piv_pred_cont[p] = 0.0

    fwd_piv_pred_int = np.round(fwd_piv_pred_cont).astype(int)

    fwd_amz_pred = fwd_piv_pred_int['Amazon']
    fwd_ebay_pred = fwd_piv_pred_int['eBay']
    fwd_web_pred = fwd_piv_pred_int['Website']
    fwd_oth_pred = fwd_piv_pred_int['Other']
    fwd_tot_pred_col = fwd_amz_pred + fwd_ebay_pred + fwd_web_pred + fwd_oth_pred

    fwd_conf_list = []
    fwd_risk_list = []
    fwd_action_list = []
    fwd_reason_list = []
    fwd_cover_list = []

    for sku in all_skus:
        r = sku_meta_df[sku_meta_df['canonical_sku'] == sku].iloc[0]
        c_pred = fwd_tot_pred_col[sku]
        stock_val = r['current_stock'] if pd.notna(r['current_stock']) else 0
        conf, risk = classify_sku_confidence_risk(stock_val, r['in_stock_flag'], r['v14'], r['v30'], r['cv_30'], c_pred)
        action = get_recommended_action(stock_val, r['in_stock_flag'], r['cv_30'], c_pred)
        reason = generate_sku_reason(stock_val, r['in_stock_flag'], r['v14'], r['v30'], r['v90'], r['cv_30'], None, c_pred)

        daily_rate = c_pred / 10.0
        if stock_val <= 0:
            cover = 0.0
        elif daily_rate <= 0:
            cover = 999.0
        else:
            cover = min(round(stock_val / daily_rate, 1), 999.0)

        fwd_conf_list.append(conf)
        fwd_risk_list.append(risk)
        fwd_action_list.append(action)
        fwd_reason_list.append(reason)
        fwd_cover_list.append(cover)

    fwd_summary_df = pd.DataFrame({
        'SKU': all_skus,
        'Product': [product_lookup[s] for s in all_skus],
        'Forecast Period': 'Sep 11–20, 2026',
        'Amazon Predicted': fwd_amz_pred.values,
        'eBay Predicted': fwd_ebay_pred.values,
        'Website Predicted': fwd_web_pred.values,
        'Other Predicted': fwd_oth_pred.values,
        'Total Predicted': fwd_tot_pred_col.values,
        'Confidence': fwd_conf_list,
        'Risk': fwd_risk_list,
        'Recommended Action': fwd_action_list,
        'Reason': fwd_reason_list
    })

    print(f"  Forward Forecast SKU Summary rows: {len(fwd_summary_df):,}")
    print(f"    Total Predicted Units (10d rounded sum): {fwd_summary_df['Total Predicted'].sum():,}")

    # =========================================================================
    # 5. BUILD WORKBOOK 1: Rimmel_Validation_Sep01_Sep10_2026.xlsx
    # =========================================================================
    print("\n[STEP 5] Generating Excel Report 1: Rimmel_Validation_Sep01_Sep10_2026.xlsx...")
    wb_v = openpyxl.Workbook()
    wb_v.remove(wb_v.active)

    # SHEET 1: SKU Validation Summary
    ws_v1 = wb_v.create_sheet('SKU Validation Summary')
    ws_v1.views.sheetView[0].showGridLines = True

    v1_headers = [
        'SKU', 'Product', 'Validation Period',
        'Amazon Actual', 'Amazon Predicted',
        'eBay Actual', 'eBay Predicted',
        'Website Actual', 'Website Predicted',
        'Other Actual', 'Other Predicted',
        'Total Actual', 'Total Predicted',
        'Variance', 'Confidence', 'Risk', 'Reason'
    ]
    v1_fills = [
        CHARCOAL_FILL, CHARCOAL_FILL, CHARCOAL_FILL,
        AMZ_FILL, AMZ_FILL,
        EBAY_FILL, EBAY_FILL,
        WEB_FILL, WEB_FILL,
        OTH_FILL, OTH_FILL,
        TOT_FILL, TOT_FILL,
        TOT_FILL, CHARCOAL_FILL, CHARCOAL_FILL, CHARCOAL_FILL
    ]
    style_header_row_custom(ws_v1, v1_headers, v1_fills)

    for row_tuple in val_summary_df[v1_headers].itertuples(index=False):
        ws_v1.append(list(row_tuple))

    ws_v1.freeze_panes = 'D2'
    ws_v1.auto_filter.ref = f"A1:Q{len(val_summary_df)+1}"

    # Format numeric columns
    for row in range(2, len(val_summary_df) + 2):
        for col_idx in [4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14]:
            c = ws_v1.cell(row=row, column=col_idx)
            c.number_format = '#,##0'
            c.alignment = Alignment(horizontal='right')
        for col_idx in [15, 16]: # Confidence, Risk
            ws_v1.cell(row=row, column=col_idx).alignment = Alignment(horizontal='center')

    auto_fit_columns_custom(ws_v1, max_cols=17, reason_col_idx=17)

    # SHEET 2: Platform Summary
    ws_v2 = wb_v.create_sheet('Platform Summary')
    ws_v2.views.sheetView[0].showGridLines = True
    ws_v2.column_dimensions['A'].width = 4
    ws_v2.column_dimensions['B'].width = 24
    ws_v2.column_dimensions['C'].width = 18
    ws_v2.column_dimensions['D'].width = 18
    ws_v2.column_dimensions['E'].width = 16
    ws_v2.column_dimensions['F'].width = 16
    ws_v2.column_dimensions['G'].width = 16

    ws_v2.cell(row=2, column=2, value="RIMMEL DEMAND FORECASTING SYSTEM").font = TITLE_FONT
    ws_v2.cell(row=3, column=2, value="Validation Platform Performance: September 1–10, 2026").font = SUBTITLE_FONT

    p2_v_headers = ['Platform', 'Actual Units', 'Predicted Units', 'Variance', 'Bias (%)', 'WAPE (%)']
    for idx, ph in enumerate(p2_v_headers, start=2):
        c = ws_v2.cell(row=5, column=idx, value=ph)
        c.font = HEADER_FONT
        c.fill = NAVY_FILL
        c.alignment = Alignment(horizontal='center')
        c.border = BORDER_ALL
    ws_v2.row_dimensions[5].height = 24

    curr_r = 6
    for plat in ['Amazon', 'eBay', 'Website', 'Other']:
        p_act = val_summary_df[f'{plat} Actual'].sum()
        p_pred = val_summary_df[f'{plat} Predicted'].sum()
        p_var = p_pred - p_act
        p_bias = ((p_pred - p_act) / p_act * 100.0) if p_act > 0 else 0.0

        # Platform WAPE
        p_sub = holdout_df[holdout_df['platform_group'] == plat]
        p_wape = (np.abs(p_sub['observed_units_sold'] - p_sub['pred_decimal']).sum() / p_sub['observed_units_sold'].sum() * 100.0) if p_sub['observed_units_sold'].sum() > 0 else 0.0

        ws_v2.append(['', plat, int(p_act), int(p_pred), int(p_var), f"{p_bias:+.2f}%", f"{p_wape:.2f}%"])
        for c_idx in range(2, 8):
            cell = ws_v2.cell(row=curr_r, column=c_idx)
            cell.border = BORDER_ALL
            if c_idx in [3, 4, 5]:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right')
            else:
                cell.alignment = Alignment(horizontal='center')
        curr_r += 1

    # Total Row
    tot_act = val_summary_df['Total Actual'].sum()
    tot_pred = val_summary_df['Total Predicted'].sum()
    tot_var = tot_pred - tot_act
    tot_bias = (tot_var / tot_act * 100.0) if tot_act > 0 else 0.0
    tot_wape = (np.abs(holdout_df['observed_units_sold'] - holdout_df['pred_decimal']).sum() / tot_act * 100.0)

    ws_v2.append(['', 'Total Portfolio', int(tot_act), int(tot_pred), int(tot_var), f"{tot_bias:+.2f}%", f"{tot_wape:.2f}%"])
    for c_idx in range(2, 8):
        cell = ws_v2.cell(row=curr_r, column=c_idx)
        cell.font = BOLD_FONT
        cell.fill = LIGHT_BLUE_FILL
        cell.border = BORDER_TOP_THIN_BOTTOM_DOUBLE
        if c_idx in [3, 4, 5]:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal='right')
        else:
            cell.alignment = Alignment(horizontal='center')

    val_excel_primary = os.path.join(reports_dir, 'Rimmel_Validation_Sep01_Sep10_2026.xlsx')
    wb_v.save(val_excel_primary)
    print(f"  Wrote primary file: {val_excel_primary}")

    # Copy to backward-compatible filename
    val_excel_compat = os.path.join(reports_dir, 'validation_report_sep_01_to_10_2026.xlsx')
    shutil.copyfile(val_excel_primary, val_excel_compat)
    print(f"  Copied to: {val_excel_compat}")

    # =========================================================================
    # 6. BUILD WORKBOOK 2: Rimmel_Forward_Forecast_Sep11_Sep20_2026.xlsx
    # =========================================================================
    print("\n[STEP 6] Generating Excel Report 2: Rimmel_Forward_Forecast_Sep11_Sep20_2026.xlsx...")
    wb_f = openpyxl.Workbook()
    wb_f.remove(wb_f.active)

    # SHEET 1: SKU Forecast Summary
    ws_f1 = wb_f.create_sheet('SKU Forecast Summary')
    ws_f1.views.sheetView[0].showGridLines = True

    f1_headers = [
        'SKU', 'Product', 'Forecast Period',
        'Amazon Predicted', 'eBay Predicted',
        'Website Predicted', 'Other Predicted',
        'Total Predicted', 'Confidence', 'Risk',
        'Recommended Action', 'Reason'
    ]
    f1_fills = [
        CHARCOAL_FILL, CHARCOAL_FILL, CHARCOAL_FILL,
        AMZ_FILL, EBAY_FILL, WEB_FILL, OTH_FILL,
        TOT_FILL, CHARCOAL_FILL, CHARCOAL_FILL,
        NAVY_FILL, CHARCOAL_FILL
    ]
    style_header_row_custom(ws_f1, f1_headers, f1_fills)

    for row_tuple in fwd_summary_df[f1_headers].itertuples(index=False):
        ws_f1.append(list(row_tuple))

    ws_f1.freeze_panes = 'D2'
    ws_f1.auto_filter.ref = f"A1:L{len(fwd_summary_df)+1}"

    for row in range(2, len(fwd_summary_df) + 2):
        for col_idx in [4, 5, 6, 7, 8]:
            c = ws_f1.cell(row=row, column=col_idx)
            c.number_format = '#,##0'
            c.alignment = Alignment(horizontal='right')
        for col_idx in [9, 10, 11]:
            ws_f1.cell(row=row, column=col_idx).alignment = Alignment(horizontal='center')

    auto_fit_columns_custom(ws_f1, max_cols=12, reason_col_idx=12)

    # SHEET 2: Platform Summary
    ws_f2 = wb_f.create_sheet('Platform Summary')
    ws_f2.views.sheetView[0].showGridLines = True
    ws_f2.column_dimensions['A'].width = 4
    ws_f2.column_dimensions['B'].width = 24
    ws_f2.column_dimensions['C'].width = 22
    ws_f2.column_dimensions['D'].width = 20

    ws_f2.cell(row=2, column=2, value="RIMMEL DEMAND FORECASTING SYSTEM").font = TITLE_FONT
    ws_f2.cell(row=3, column=2, value="Forward Forecast Platform Breakdown: September 11–20, 2026").font = SUBTITLE_FONT

    p2_f_headers = ['Platform', '10-Day Predicted Units', 'Channel Share (%)']
    for idx, ph in enumerate(p2_f_headers, start=2):
        c = ws_f2.cell(row=5, column=idx, value=ph)
        c.font = HEADER_FONT
        c.fill = NAVY_FILL
        c.alignment = Alignment(horizontal='center')
        c.border = BORDER_ALL
    ws_f2.row_dimensions[5].height = 24

    f_tot_units = fwd_summary_df['Total Predicted'].sum()
    curr_r = 6
    for plat in ['Amazon', 'eBay', 'Website', 'Other']:
        p_pred = fwd_summary_df[f'{plat} Predicted'].sum()
        p_share = (p_pred / f_tot_units * 100.0) if f_tot_units > 0 else 0.0

        ws_f2.append(['', plat, int(p_pred), f"{p_share:.1f}%"])
        for c_idx in range(2, 5):
            cell = ws_f2.cell(row=curr_r, column=c_idx)
            cell.border = BORDER_ALL
            if c_idx == 3:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right')
            else:
                cell.alignment = Alignment(horizontal='center')
        curr_r += 1

    ws_f2.append(['', 'Total Portfolio', int(f_tot_units), "100.0%"])
    for c_idx in range(2, 5):
        cell = ws_f2.cell(row=curr_r, column=c_idx)
        cell.font = BOLD_FONT
        cell.fill = LIGHT_BLUE_FILL
        cell.border = BORDER_TOP_THIN_BOTTOM_DOUBLE
        if c_idx == 3:
            cell.number_format = '#,##0'
            cell.alignment = Alignment(horizontal='right')
        else:
            cell.alignment = Alignment(horizontal='center')

    # SHEET 3: Inventory Actions
    ws_f3 = wb_f.create_sheet('Inventory Actions')
    ws_f3.views.sheetView[0].showGridLines = True

    f3_headers = [
        'SKU', 'Product', '10-Day Forecast',
        'Current Shared Warehouse Stock', 'Days of Cover',
        'Inventory Risk Status', 'Recommended Action'
    ]
    f3_fills = [
        CHARCOAL_FILL, CHARCOAL_FILL, TOT_FILL,
        NAVY_FILL, NAVY_FILL, CHARCOAL_FILL, NAVY_FILL
    ]
    style_header_row_custom(ws_f3, f3_headers, f3_fills)

    for sku in all_skus:
        r = sku_meta_df[sku_meta_df['canonical_sku'] == sku].iloc[0]
        c_pred = fwd_tot_pred_col[sku]
        stock_val = int(r['current_stock']) if pd.notna(r['current_stock']) else 0
        daily_rate = c_pred / 10.0
        if stock_val <= 0: cover = 0.0
        elif daily_rate <= 0: cover = 999.0
        else: cover = min(round(stock_val / daily_rate, 1), 999.0)

        _, risk = classify_sku_confidence_risk(stock_val, r['in_stock_flag'], r['v14'], r['v30'], r['cv_30'], c_pred)
        action = get_recommended_action(stock_val, r['in_stock_flag'], r['cv_30'], c_pred)

        ws_f3.append([
            sku,
            product_lookup[sku],
            int(c_pred),
            int(stock_val),
            cover,
            risk,
            action
        ])

    ws_f3.freeze_panes = 'C2'
    ws_f3.auto_filter.ref = f"A1:G{len(all_skus)+1}"

    for row in range(2, len(all_skus) + 2):
        ws_f3.cell(row=row, column=3).number_format = '#,##0'
        ws_f3.cell(row=row, column=3).alignment = Alignment(horizontal='right')
        ws_f3.cell(row=row, column=4).number_format = '#,##0'
        ws_f3.cell(row=row, column=4).alignment = Alignment(horizontal='right')
        ws_f3.cell(row=row, column=5).number_format = '#,##0.0'
        ws_f3.cell(row=row, column=5).alignment = Alignment(horizontal='right')
        ws_f3.cell(row=row, column=6).alignment = Alignment(horizontal='center')
        ws_f3.cell(row=row, column=7).alignment = Alignment(horizontal='center')

    auto_fit_columns_custom(ws_f3, max_cols=7)

    fwd_excel_primary = os.path.join(reports_dir, 'Rimmel_Forward_Forecast_Sep11_Sep20_2026.xlsx')
    wb_f.save(fwd_excel_primary)
    print(f"  Wrote primary file: {fwd_excel_primary}")

    # Copy to backward-compatible filename
    fwd_excel_compat = os.path.join(reports_dir, 'production_forecast_sep_11_to_20_2026.xlsx')
    shutil.copyfile(fwd_excel_primary, fwd_excel_compat)
    print(f"  Copied to: {fwd_excel_compat}")

    # 7. CACHE EXPORTS FOR DASHBOARDS & CSV
    val_summary_df.to_csv(os.path.join(processed_dir, 'dashboard_validation_sku_summary.csv'), index=False)
    fwd_summary_df.to_csv(os.path.join(processed_dir, 'dashboard_forecast_sku_summary.csv'), index=False)

    print("\n" + "=" * 80)
    print("REPORT GENERATION COMPLETE")
    print(f"Primary Validation Report: {val_excel_primary}")
    print(f"Primary Forward Forecast:  {fwd_excel_primary}")
    print("=" * 80)

if __name__ == '__main__':
    main()
